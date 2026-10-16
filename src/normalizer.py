"""Excel Semantic Normalizer: Unwraps merged cells and flattens hierarchical headers."""

import re
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl
import pandas as pd


def clean_cell_value(val: Any) -> Any:
    """Cleans cell value, stripping excessive whitespace and unwrapping hyphenated linebreaks."""
    if isinstance(val, str):
        val = re.sub(r'-\s*\n\s*', '', val)
        val = re.sub(r'\s*\n\s*', ' ', val)
        val = re.sub(r'\s+', ' ', val).strip()
    return val


def is_num_cell(x: Any) -> bool:
    """Checks if a cell represents a numeric value (int, float, currency, percent, comma-formatted)."""
    if x is None:
        return False
    if isinstance(x, (int, float)):
        return True
    s = str(x).strip()
    # Strip currency symbols, percentages, and thousands commas
    s_clean = re.sub(r'[\$,%€£¥₫]', '', s).strip()
    # Handle negative in parentheses e.g. (1.5) -> -1.5
    if s_clean.startswith('(') and s_clean.endswith(')'):
        s_clean = '-' + s_clean[1:-1].strip()
    return bool(re.match(r'^-?[\d,]+(\.\d+)?$', s_clean))


def is_title_or_meta_row(row_vals: List[Any], total_cols: int) -> bool:
    """Identifies metadata, table titles, subtitles, or blank rows."""
    non_empty = [str(x).strip() for x in row_vals if pd.notna(x) and str(x).strip() not in ['', 'nan', 'None']]
    if not non_empty:
        return True  # Blank spacer row

    # If all non-empty cells across the row are identical and table has >2 columns (full-width merged title)
    if len(set(non_empty)) == 1 and total_cols > 2:
        return True

    # Standard BLS/Census/StatCan/SEC metadata prefixes
    first_str = non_empty[0].lower()
    meta_prefixes = [
        'table ', 'table:', 'tableau ', 'household data',
        '[numbers in', '[percent', '[in thousands', '[in millions',
        '(dollars in', '(in millions', '(in thousands',
        'thu nhập tiền lương', '25至54岁'
    ]
    if any(first_str.startswith(k) for k in meta_prefixes):
        return True

    return False


def make_unique_columns(cols: List[str]) -> List[str]:
    """Ensures all column names in the list are strictly unique without collision."""
    seen: Dict[str, int] = {}
    unique_cols = []
    for c in cols:
        name = c.strip() if (c and str(c).strip()) else 'Col'
        if name not in seen:
            seen[name] = 1
            unique_cols.append(name)
        else:
            cand = f"{name}_{seen[name]}"
            while cand in seen or cand in cols:
                seen[name] += 1
                cand = f"{name}_{seen[name]}"
            seen[name] += 1
            seen[cand] = 1
            unique_cols.append(cand)
    return unique_cols


def normalize_worksheet(ws: openpyxl.worksheet.worksheet.Worksheet) -> Optional[pd.DataFrame]:
    """Converts a single openpyxl worksheet into a cleanly structured pandas DataFrame with unmerged cells."""
    if ws.max_row is None or ws.max_column is None:
        return None

    max_r = ws.max_row
    max_c = ws.max_column

    # Step 1: Read all raw cells into a 2D matrix
    grid: List[List[Any]] = [[None for _ in range(max_c)] for _ in range(max_r)]
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            val = ws.cell(r, c).value
            grid[r - 1][c - 1] = clean_cell_value(val)

    # Step 2: Propagate merged cell values across the entire merged range
    for merged_range in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left_val = grid[min_row - 1][min_col - 1]
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                grid[r - 1][c - 1] = top_left_val

    raw_df = pd.DataFrame(grid).dropna(how="all").dropna(how="all", axis=1)
    if raw_df.empty:
        return None

    raw_df.reset_index(drop=True, inplace=True)
    return clean_and_headerize(raw_df)


def clean_and_headerize(df: pd.DataFrame) -> pd.DataFrame:
    """Removes metadata rows, merges multi-level headers into unique composite column names, and propagates section headers."""
    total_cols = df.shape[1]

    # Step 1: Strip leading title, metadata, and blank rows (up to first 6 rows)
    start_r = 0
    while start_r < min(6, len(df)):
        if is_title_or_meta_row(list(df.iloc[start_r]), total_cols):
            start_r += 1
        else:
            break
    sub_df = df.iloc[start_r:].reset_index(drop=True)
    if sub_df.empty:
        return df

    # Step 2: Detect boundary between column headers and data/section rows
    header_end = 1
    for r in range(min(8, len(sub_df))):
        row_vals = sub_df.iloc[r]
        nums = [x for x in row_vals if is_num_cell(x)]

        # Filter out standalone calendar years (1900-2100) often appearing in headers (e.g. '2022', '2023')
        non_year_nums = []
        for x in nums:
            try:
                v = float(re.sub(r'[^\d.-]', '', str(x)))
                if v < 1900 or v > 2100:
                    non_year_nums.append(v)
            except Exception:
                pass

        # Data row detected if multiple columns contain actual numeric measurements
        if len(non_year_nums) >= 2 or (total_cols <= 3 and len(non_year_nums) >= 1):
            header_end = r
            break

        # Section header row detected: only col 0 or col 1 has text, while all data columns (col 2+) are empty
        other_vals = [x for x in row_vals.iloc[2:] if pd.notna(x) and str(x).strip() not in ['', 'nan', 'None']]
        first_vals = [x for x in row_vals.iloc[:2] if pd.notna(x) and str(x).strip() not in ['', 'nan', 'None']]
        if r > 0 and len(first_vals) > 0 and len(other_vals) == 0:
            header_end = r
            break

    # Step 3: Combine hierarchical headers into unique composite column names
    cols = []
    for c in range(sub_df.shape[1]):
        parts = []
        for hr in range(header_end):
            v = str(sub_df.iloc[hr, c]).strip()
            if v and v not in ['nan', 'None', ''] and v not in parts:
                parts.append(v)
        cols.append(' | '.join(parts) if parts else f'Col_{c}')
    cols = make_unique_columns(cols)

    data_df = sub_df.iloc[header_end:].reset_index(drop=True)
    data_df.columns = cols
    data_df = data_df.dropna(how="all").reset_index(drop=True)

    # Step 4: Propagate section headers in column 0
    current_section = ''
    clean_rows = []
    for idx in range(len(data_df)):
        row = data_df.iloc[idx].copy()
        val0 = str(row.iloc[0]).strip()
        other_vals = [x for x in row.iloc[1:] if pd.notna(x) and str(x).strip() not in ['', 'nan', 'None']]

        # Section header row has non-empty col 0 but no numeric data in other columns
        if len(other_vals) == 0 and val0 not in ['', 'nan', 'None'] and not any(k in val0 for k in ['Footnotes:', 'NOTE:', 'Source:']):
            current_section = val0
        else:
            if current_section and not val0.startswith(current_section) and not any(k in val0 for k in ['Footnotes:', 'NOTE:', 'Source:']):
                row.iloc[0] = f'{current_section} | {val0}'
            clean_rows.append(row)

    if clean_rows:
        return pd.DataFrame(clean_rows).reset_index(drop=True)
    return data_df


def load_normalized_workbook(file_path: str | Path) -> Dict[str, pd.DataFrame]:
    """Loads an .xlsx file, resolves all merged cells, and converts each sheet to a normalized DataFrame."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets_data: Dict[str, pd.DataFrame] = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        df = normalize_worksheet(ws)
        if df is not None and not df.empty:
            sheets_data[sheet_name] = df

    return sheets_data
