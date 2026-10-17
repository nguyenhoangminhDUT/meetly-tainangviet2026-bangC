"""Comprehensive Test Suite for Meetly Spreadsheet QA Pipeline."""

import sys
from pathlib import Path
import pandas as pd

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.normalizer import load_normalized_workbook, make_unique_columns, clean_cell_value, is_num_cell, extract_worksheet_tables
from src.sanitizer import sanitize_answer, strip_thousands_separators
from src.schema_builder import build_schema_summary, extract_keywords
from src.solver import execute_code_safely, is_empty_or_meaningless, get_column_samples_hint
from src.evaluator import compute_levenshtein, normalize_for_eval, compute_relative_numeric_error


def test_unique_columns():
    cols = ["col", "col", "col_1", "col"]
    uniq = make_unique_columns(cols)
    assert len(uniq) == len(set(uniq))
    assert uniq[0] == "col"
    assert uniq[2] == "col_1"


def test_is_num_cell():
    assert is_num_cell(42) is True
    assert is_num_cell(3.14) is True
    assert is_num_cell("58,176") is True
    assert is_num_cell("$1,402") is True
    assert is_num_cell("10%") is True
    assert is_num_cell("(1.5)") is True
    assert is_num_cell("Not a number") is False
    assert is_num_cell("..") is False
    assert is_num_cell("-") is False


def test_strip_thousands_separators():
    assert strip_thousands_separators("7,668") == "7668"
    assert strip_thousands_separators("1,094,431") == "1094431"
    assert strip_thousands_separators("2020, 9743") == "2020, 9743"
    assert strip_thousands_separators("34,282, 118,705") == "34282, 118705"


def test_multilingual_sanitizer():
    # Exactly one decimal place
    assert sanitize_answer("4.71", "Answer with only the number using exactly one decimal place.") == "4.7"
    assert sanitize_answer("50", "Chỉ trả lời bằng số với đúng một chữ số thập phân.") == "50.0"
    assert sanitize_answer("10.42", "Répondez uniquement par le nombre avec une décimale, sans le symbole %.") == "10.4"
    assert sanitize_answer("19.12", "只回答数字，保留一位小数。") == "19.1"

    # Exactly two decimal places
    assert sanitize_answer("1.2", "Answer with only the number using exactly two decimal places.") == "1.20"
    assert sanitize_answer("14.2", "소수점 둘째 자리까지, % 기호 없이 숫자만 답하십시오.") == "14.20"
    assert sanitize_answer("40.6 and 37.6", "Answer as `men, women`, using exactly two decimal places and separating the values with a comma and one space.") == "40.60, 37.60"

    # Exactly three decimal places
    assert sanitize_answer("29.4921", "Answer with only the number using exactly three decimal places.") == "29.492"
    assert sanitize_answer("0.8709", "rounded to three decimal places") == "0.871"

    # Only the Year
    assert sanitize_answer("The year was 1954", "Identify the year when... Answer with only the year.") == "1954"
    assert sanitize_answer("2013", "in which year did the... Answer with only the year.") == "2013"

    # Integers without commas
    assert sanitize_answer("27,124 dollars", "Chỉ trả lời bằng số nguyên, không có dấu phẩy.") == "27124"
    assert sanitize_answer("2020 and 9743", "Answer as `year, decrease`, separated by a comma and one space, as integers without commas or units.") == "2020, 9743"
    assert sanitize_answer("1955, 62,170", "Answer as `year, total employed`, separated by a comma and one space, with the total as an integer without commas or units.") == "1955, 62170"

    # Yes / No
    assert sanitize_answer("yes", "Answer with only Yes or No.") == "Yes"
    assert sanitize_answer("no", "Answer with only Yes or No.") == "No"

    # Only last name (Multicultural: English last name vs Vietnamese surname)
    assert sanitize_answer("Minh Thi Ngo", "Answer with only the last name.") == "Ngo"
    assert sanitize_answer("John Smith", "Answer with only the last name.") == "Smith"
    assert sanitize_answer("Nguyễn Văn An", "Chỉ họ của người đó:") == "Nguyễn"
    assert sanitize_answer("Tran Minh Tu", "Chỉ ghi họ:") == "Tran"
    assert sanitize_answer("Ngo", "Answer with only the last name.") == "Ngo"

    # Tuple / List preservation
    assert sanitize_answer((2020, 9743), "as integers without commas or units") == "2020, 9743"
    assert sanitize_answer(["table_1", 1307], "as integers without commas or units") == "table_1, 1307"


def test_schema_builder_keywords():
    q = "What is the unadjusted percent change for Housing and Food?"
    kws = extract_keywords(q)
    assert "housing" in kws
    assert "food" in kws
    assert "unadjusted" in kws


def test_sandboxed_code_execution():
    code = """
def solve(dfs):
    df = dfs['test']
    return df['A'].sum()
"""
    dfs = {'test': pd.DataFrame({'A': [10, 20, 30]})}
    res, err = execute_code_safely(code, dfs)
    assert err is None
    assert res == 60


def test_all_public_workbooks_normalization():
    for d in ["data/data_public_1/xlsx_files", "data/Public data 2/xlsx_files"]:
        p = Path(d)
        if not p.exists():
            continue
        for f in sorted(p.glob("*.xlsx")):
            if f.name.startswith("~$"):
                continue
            dfs = load_normalized_workbook(f)
            assert len(dfs) > 0, f"No sheets returned for {f.name}"
            for sname, df in dfs.items():
                assert not df.empty, f"Empty DataFrame for {f.name}::{sname}"
                # Guarantee column uniqueness
                assert len(df.columns) == len(set(df.columns)), f"Duplicate columns in {f.name}::{sname}"


def test_reflection_helpers():
    import numpy as np
    assert is_empty_or_meaningless(None) is True
    assert is_empty_or_meaningless("") is True
    assert is_empty_or_meaningless("   ") is True
    assert is_empty_or_meaningless(np.nan) is True
    assert is_empty_or_meaningless("nan") is True
    assert is_empty_or_meaningless("None") is True
    assert is_empty_or_meaningless([]) is True
    assert is_empty_or_meaningless({}) is True
    assert is_empty_or_meaningless(pd.DataFrame()) is True
    assert is_empty_or_meaningless("Valid Answer") is False
    assert is_empty_or_meaningless(42) is False
    assert is_empty_or_meaningless(0) is False

    test_dfs = {
        "Sheet1": pd.DataFrame({
            "Category": ["Agriculture", "Manufacturing", "Services"],
            "Value": [100, 200, 300]
        })
    }
    hints = get_column_samples_hint(test_dfs)
    assert "Category" in hints
    assert "Agriculture" in hints


def test_evaluator_diagnostics():
    assert compute_levenshtein("kitten", "sitting") == 3
    assert compute_levenshtein("equal", "equal") == 0
    assert normalize_for_eval("11.0") == "11"
    assert normalize_for_eval("$4.70") == "4.7"
    assert normalize_for_eval("Yes.") == "yes"
    assert compute_relative_numeric_error("100", "100") == 0.0
    assert abs(compute_relative_numeric_error("105", "100") - 0.05) < 1e-5


def test_multi_table_extraction():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "MultiSheet"
    
    # Table 1: rows 1-3
    ws.cell(1, 1, "Item")
    ws.cell(1, 2, "Sales")
    ws.cell(2, 1, "Apples")
    ws.cell(2, 2, 100)
    ws.cell(3, 1, "Oranges")
    ws.cell(3, 2, 200)

    # Empty spacer rows 4, 5
    # Table 2: rows 6-8
    ws.cell(6, 1, "Region")
    ws.cell(6, 2, "Profit")
    ws.cell(7, 1, "North")
    ws.cell(7, 2, 50)
    ws.cell(8, 1, "South")
    ws.cell(8, 2, 80)

    tables = extract_worksheet_tables(ws, sheet_name="MultiSheet")
    assert len(tables) == 2
    assert "MultiSheet" in tables
    assert "MultiSheet_sub2" in tables
    assert "Sales" in tables["MultiSheet"].columns
    assert "Profit" in tables["MultiSheet_sub2"].columns


if __name__ == "__main__":
    print("Running test suite...")
    test_unique_columns()
    test_is_num_cell()
    test_strip_thousands_separators()
    test_multilingual_sanitizer()
    test_schema_builder_keywords()
    test_sandboxed_code_execution()
    test_all_public_workbooks_normalization()
    test_reflection_helpers()
    test_evaluator_diagnostics()
    test_multi_table_extraction()
    print(">>> ALL 10 TEST SUITES PASSED FLAWLESSLY! <<<")
